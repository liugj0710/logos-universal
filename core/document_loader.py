"""
文档解析 + 智能分段
支持：DOCX, PDF, TXT/Markdown, Excel(.xlsx)
"""
import os
import re
import hashlib
from typing import List, Dict, Any

try:
    import docx
except ImportError:
    docx = None

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


def get_file_hash(filepath: str) -> str:
    """计算文件 SHA256（前 16 位）"""
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(8192):
            h.update(chunk)
    return h.hexdigest()[:16]


def split_chunks(text: str, max_len: int = 600, overlap: int = 50) -> List[str]:
    """
    按语义切分文本块。
    优先按段落切，段落过长再按句子切。
    """
    if not text:
        return []
    
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks = []
    current = ""
    
    for para in paragraphs:
        if len(para) > max_len * 1.5:
            # 超长段落按句子切
            sentences = re.split(r'([。！？；\n])', para)
            sentences = [s for s in sentences if s.strip()]
            for s in sentences:
                if len(current) + len(s) > max_len and current:
                    chunks.append(current.strip())
                    current = current[-overlap:] if overlap < len(current) else ""
                current += s
        else:
            if len(current) + len(para) > max_len and current:
                chunks.append(current.strip())
                current = current[-overlap:] if overlap < len(current) else ""
            current += para + "\n\n"
    
    if current.strip():
        chunks.append(current.strip())
    
    return [c for c in chunks if len(c) >= 50]


def parse_docx(filepath: str) -> List[Dict[str, Any]]:
    """解析 DOCX，保留标题层级"""
    if docx is None:
        raise ImportError("python-docx not installed")
    
    doc = docx.Document(filepath)
    chunks = []
    current_heading = ""
    full_text = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        
        if para.style.name.startswith("Heading"):
            if full_text:
                content = "\n\n".join(full_text)
                for chunk in split_chunks(content):
                    chunks.append({
                        "text": chunk,
                        "title": current_heading,
                        "page": None,
                        "sheet": None,
                    })
                full_text = []
            current_heading = text
        else:
            full_text.append(text)
    
    if full_text:
        content = "\n\n".join(full_text)
        for chunk in split_chunks(content):
            chunks.append({
                "text": chunk,
                "title": current_heading,
                "page": None,
                "sheet": None,
            })
    
    return chunks


def parse_pdf(filepath: str) -> List[Dict[str, Any]]:
    """解析 PDF，按页提取"""
    if fitz is None:
        raise ImportError("PyMuPDF not installed")
    
    doc = fitz.open(filepath)
    chunks = []
    
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text = page.get_text()
        if not text.strip():
            continue
        
        page_chunks = split_chunks(text, max_len=800, overlap=50)
        for chunk in page_chunks:
            chunks.append({
                "text": chunk,
                "title": f"第{page_num + 1}页",
                "page": page_num + 1,
                "sheet": None,
            })
    
    doc.close()
    return chunks


def parse_txt(filepath: str) -> List[Dict[str, Any]]:
    """解析 TXT/Markdown"""
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    
    chunks = split_chunks(text)
    return [{
        "text": c,
        "title": os.path.basename(filepath),
        "page": None,
        "sheet": None,
    } for c in chunks]


def parse_excel(filepath: str) -> List[Dict[str, Any]]:
    """解析 Excel，每 Sheet 每 50 行一个 chunk"""
    if openpyxl is None:
        raise ImportError("openpyxl not installed")
    
    wb = openpyxl.load_workbook(filepath, data_only=True)
    chunks = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        headers = []
        
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(c) if c is not None else "" for c in row]
                continue
            
            row_data = [str(c) if c is not None else "" for c in row]
            if any(row_data):
                row_dict = dict(zip(headers, row_data))
                rows.append(str(row_dict))
            
            if len(rows) >= 50:
                content = f"Sheet: {sheet_name}\n" + "\n".join(rows)
                chunks.append({
                    "text": content,
                    "title": f"{os.path.basename(filepath)} - {sheet_name}",
                    "page": None,
                    "sheet": sheet_name,
                })
                rows = []
        
        if rows:
            content = f"Sheet: {sheet_name}\n" + "\n".join(rows)
            chunks.append({
                "text": content,
                "title": f"{os.path.basename(filepath)} - {sheet_name}",
                "page": None,
                "sheet": sheet_name,
            })
    
    wb.close()
    return chunks


def load_document(filepath: str) -> List[Dict[str, Any]]:
    """
    统一入口，根据后缀分发解析器。
    返回: [{"text": "...", "title": "...", "page": ..., "sheet": ...}, ...]
    """
    ext = os.path.splitext(filepath)[1].lower()
    
    if ext == ".docx":
        return parse_docx(filepath)
    elif ext == ".pdf":
        return parse_pdf(filepath)
    elif ext in (".txt", ".md", ".markdown"):
        return parse_txt(filepath)
    elif ext in (".xlsx", ".xls"):
        return parse_excel(filepath)
    else:
        return []